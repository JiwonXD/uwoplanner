"""Import the supplied workbook's saved values; never execute its formulas.

Usage: python scripts/import_navigators.py SOURCE.xlsx [--output data]
Requires openpyxl for reading only. Outputs SQLite and browser-friendly JSON.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sqlite3

import openpyxl


STAT_NAMES = ["포격술", "충파술", "지원술", "백병술", "박물학", "심미학",
              "척후법", "보급법", "구매 전략", "판매 전략", "협상 전략", "교환 전략"]


def stable_id(kind, name):
    return kind + "_" + hashlib.sha256(name.encode("utf-8")).hexdigest()[:20]


def parse_ability(raw):
    match = re.fullmatch(r"(.+?)\s+LV\s*(\d+)", raw.strip(), re.IGNORECASE)
    return (match[1].strip(), int(match[2])) if match else (raw.strip(), None)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data"))
    args = parser.parse_args()
    values = openpyxl.load_workbook(args.source, data_only=True)
    formulas = openpyxl.load_workbook(args.source, data_only=False)
    source = {
        "filename": args.source.name,
        "sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
        "version_label": "v260708" if "v260708" in args.source.name else None,
        "author_label": "아스트라이오스",
        "value_mode": "xlsx_saved_values_no_recalculation",
        "stat_basis": None,
        "stat_basis_status": "성장·장비·레벨 기준 확인 필요",
    }
    # These curated index sheets cover only some abilities. Record evidence,
    # not a guessed category for every ability in the master sheet.
    evidence = {}
    for sheet in values:
        if "효과" not in sheet.title:
            continue
        category = "모험" if sheet.title.startswith("모험") else "교역" if sheet.title.startswith("교역") else "전투"
        for row in sheet:
            for cell in row[:-1]:
                adjacent = sheet.cell(cell.row, cell.column + 1).value
                if isinstance(cell.value, str) and re.fullmatch(r"\(\d+명\)", str(adjacent)):
                    name, _ = parse_ability(cell.value)
                    evidence.setdefault(name, []).append({"category": category, "sheet": sheet.title, "cell": cell.coordinate})

    sheet = values["항해사"]
    expected_headers = {3: "이름", 4: "타입", 5: "직업", 18: "Lv 10 직업 효과", 30: "교전 기술 LV 1", 32: "인연 연대기 추가 기술"}
    for col, label in expected_headers.items():
        if sheet.cell(1, col).value != label:
            raise ValueError(f"Unsupported layout: {sheet.cell(1, col).coordinate}")

    navigators, abilities, grants = [], {}, []
    names = set()
    errors = []
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row, 3).value
        if name is None:
            continue
        if name in names:
            raise ValueError(f"Duplicate navigator: {name}")
        names.add(name)
        navigator_id = stable_id("navigator", name)
        stats = {}
        for col, stat in enumerate(STAT_NAMES, 6):
            cell = sheet.cell(row, col)
            if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
                errors.append({"sheet": sheet.title, "cell": cell.coordinate, "issue": "missing_or_non_numeric_stat"})
            stats[stat] = cell.value
        navigators.append({"id": navigator_id, "name": name,
                           "grade": sheet.cell(row, 1).value,
                           "type": sheet.cell(row, 4).value,
                           "job": sheet.cell(row, 5).value,
                           "stats": stats, "source_sheet": sheet.title, "source_row": row})
        for col in range(18, 33):
            cell = sheet.cell(row, col)
            if cell.value is None:
                continue
            if cell.data_type == "e" or not isinstance(cell.value, str):
                raise ValueError(f"Invalid ability at {cell.coordinate}: {cell.value}")
            kind = "effect" if col < 30 else "skill"
            ability_name, explicit_level = parse_ability(cell.value)
            ability_id = stable_id(kind, ability_name)
            supporting = evidence.get(ability_name, [])
            categories = sorted({item["category"] for item in supporting})
            abilities[ability_id] = {
                "id": ability_id, "name": ability_name, "kind": kind,
                "category": "전투" if kind == "skill" else categories[0] if len(categories) == 1 else None,
                "category_evidence": supporting,
                "scope": "ship" if kind == "skill" else None,
                "scope_basis": "user_statement" if kind == "skill" else "requires_per_effect_review",
                "stacking_rule": None, "max_level": None,
            }
            header = sheet.cell(1, col).value
            level_match = re.search(r"LV\s*(\d+)", header, re.IGNORECASE)
            unlock_level = int(level_match[1]) if level_match else None
            origin = "potential" if col in (26, 27) else "relationship" if col in (28, 29, 32) else "job" if col in (18, 19, 22, 23) else "character" if col < 30 else "level_skill"
            grants.append({"navigator_id": navigator_id, "ability_id": ability_id,
                           "source_column": openpyxl.utils.get_column_letter(col),
                           "source_cell": cell.coordinate, "source_header": header,
                           "origin": origin, "unlock_level": unlock_level,
                           "explicit_level": explicit_level,
                           "stack_contribution": None, "raw_text": cell.value})

    if errors:
        raise ValueError(json.dumps(errors, ensure_ascii=False))
    payload = {"schema_version": 1, "source": source,
               "navigators": navigators, "abilities": list(abilities.values()), "grants": grants}
    report = {"navigator_count": len(navigators),
              "grades": dict(Counter(n["grade"] for n in navigators)),
              "stat_values": len(navigators) * len(STAT_NAMES),
              "ability_counts": dict(Counter(a["kind"] for a in abilities.values())),
              "grant_counts": dict(Counter(abilities[g["ability_id"]]["kind"] for g in grants)),
              "explicit_level_grants": sum(g["explicit_level"] is not None for g in grants),
              "uncategorized_effects": sum(a["kind"] == "effect" and a["category"] is None for a in abilities.values()),
              "unresolved_effect_scopes": sum(a["kind"] == "effect" for a in abilities.values()),
              "duplicate_names": 0, "invalid_stats": 0,
              "limitations": ["파일에 저장된 값의 스냅샷이며 최신 게임 데이터 여부는 검증하지 않음",
                              "스탯의 성장·장비·레벨 기준 미확인",
                              "효과 분류는 효과별 목록에 존재하는 이름만 근거와 함께 채움",
                              "LV2는 원문 표기 레벨로 보존하며 중첩 기여량으로 확정하지 않음",
                              "효과 적용 범위, 중첩 공식, 상한 및 활성화 조건 추가 확인 필요",
                              "순위표의 주관적 점수는 스탯에 포함하지 않음"]}
    args.output.mkdir(parents=True, exist_ok=True)
    db_path = args.output / "navigators.sqlite"
    staging = args.output / "navigators.build.sqlite"
    if staging.exists():
        raise FileExistsError(f"Previous unfinished import exists: {staging}")
    db = sqlite3.connect(staging)
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
        CREATE TABLE source (id INTEGER PRIMARY KEY, metadata_json TEXT NOT NULL);
        CREATE TABLE navigator (id TEXT PRIMARY KEY, name TEXT NOT NULL UNIQUE,
          grade TEXT NOT NULL, type TEXT NOT NULL, job TEXT NOT NULL,
          source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL);
        CREATE TABLE navigator_stat (navigator_id TEXT REFERENCES navigator(id),
          name TEXT NOT NULL, value REAL NOT NULL, basis TEXT,
          PRIMARY KEY(navigator_id, name));
        CREATE TABLE ability (id TEXT PRIMARY KEY, name TEXT NOT NULL,
          kind TEXT NOT NULL CHECK(kind IN ('effect','skill')), category TEXT,
          scope TEXT, scope_basis TEXT NOT NULL, stacking_rule TEXT, max_level INTEGER,
          category_evidence_json TEXT NOT NULL, UNIQUE(kind,name));
        CREATE TABLE navigator_ability (navigator_id TEXT REFERENCES navigator(id),
          ability_id TEXT REFERENCES ability(id), source_column TEXT NOT NULL,
          source_cell TEXT NOT NULL, source_header TEXT NOT NULL, origin TEXT NOT NULL,
          unlock_level INTEGER, explicit_level INTEGER, stack_contribution REAL,
          raw_text TEXT NOT NULL, PRIMARY KEY(navigator_id,source_column));
        CREATE INDEX ability_holders ON navigator_ability(ability_id);
        CREATE TABLE source_cell (sheet TEXT NOT NULL, cell TEXT NOT NULL,
          value_json TEXT NOT NULL, formula TEXT, PRIMARY KEY(sheet,cell));
    """)
    db.execute("INSERT INTO source VALUES(1,?)", (json.dumps(source, ensure_ascii=False),))
    for n in navigators:
        db.execute("INSERT INTO navigator VALUES(?,?,?,?,?,?,?)", tuple(n[k] for k in ("id", "name", "grade", "type", "job", "source_sheet", "source_row")))
        db.executemany("INSERT INTO navigator_stat VALUES(?,?,?,NULL)", [(n["id"], k, v) for k, v in n["stats"].items()])
    for a in abilities.values():
        db.execute("INSERT INTO ability VALUES(?,?,?,?,?,?,?,?,?)", tuple(a[k] for k in ("id", "name", "kind", "category", "scope", "scope_basis", "stacking_rule", "max_level")) + (json.dumps(a["category_evidence"], ensure_ascii=False),))
    for g in grants:
        db.execute("INSERT INTO navigator_ability VALUES(?,?,?,?,?,?,?,?,?,?)", tuple(g.values()))
    for s in values:
        for row in s:
            for cell in row:
                original = formulas[s.title][cell.coordinate]
                if cell.value is not None or original.data_type == "f":
                    formula = original.value if original.data_type == "f" else None
                    if formula is not None and not isinstance(formula, str):
                        formula = json.dumps({"text": formula.text, "ref": formula.ref}, ensure_ascii=False)
                    db.execute("INSERT INTO source_cell VALUES(?,?,?,?)", (s.title, cell.coordinate, json.dumps(cell.value, ensure_ascii=False, default=str), formula))
    assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert not db.execute("PRAGMA foreign_key_check").fetchall()
    assert db.execute("SELECT count(*) FROM navigator_stat").fetchone()[0] == report["stat_values"]
    assert db.execute("SELECT count(*) FROM navigator_ability").fetchone()[0] == len(grants)
    db.commit()
    db.close()
    staging.replace(db_path)
    for filename, content in (("navigators.json", payload), ("import-report.json", report)):
        (args.output / filename).write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
